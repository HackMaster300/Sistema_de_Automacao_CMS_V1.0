"""Fonte de dados a processar, com checkpoint.

Os identificadores a processar (e o progresso da automação) vivem num
CSV ou Excel — nunca no código, e nunca no appsettings.yaml (que é só
para configuração da aplicação/UI).

Cada linha tem três colunas:
    identificador   — o valor a inserir no sistema legado
    estado          — "pendente" | "concluido" | "erro"
    processado_em   — timestamp ISO da última tentativa (informativo)

Isto permite RETOMAR uma execução interrompida: ao (re)correr a
automação, apenas as linhas com estado "pendente" são processadas — as
já marcadas "concluido" são ignoradas. O estado é gravado no ficheiro
imediatamente a seguir a cada registo processado (não só no fim), para
que uma falha a meio não perca o progresso já feito.
"""
from __future__ import annotations

import csv
import os
import tempfile
import time
from abc import ABC, abstractmethod
from dataclasses import dataclass
from enum import Enum
from pathlib import Path

COLUMN_IDENTIFICADOR = "identificador"
COLUMN_ESTADO = "estado"
COLUMN_PROCESSADO_EM = "processado_em"


class RecordStatus(str, Enum):
    PENDENTE = "pendente"
    CONCLUIDO = "concluido"
    ERRO = "erro"

    @classmethod
    def from_text(cls, texto: str | None) -> "RecordStatus":
        texto = (texto or "").strip().lower()
        for status in cls:
            if status.value == texto:
                return status
        return cls.PENDENTE


@dataclass
class DataRecord:
    identificador: str
    estado: RecordStatus
    processado_em: str | None = None


class CheckpointDataSource(ABC):
    """Interface comum a qualquer formato de ficheiro de dados."""

    @abstractmethod
    def load_all(self) -> list[DataRecord]:
        """Devolve todas as linhas, na ordem do ficheiro."""

    @abstractmethod
    def mark(self, identificador: str, estado: RecordStatus) -> None:
        """Atualiza o estado de uma linha e grava imediatamente (checkpoint)."""

    def pending(self) -> list[DataRecord]:
        return [r for r in self.load_all() if r.estado is RecordStatus.PENDENTE]

    def summary(self) -> dict[str, int]:
        counts = {status.value: 0 for status in RecordStatus}
        for record in self.load_all():
            counts[record.estado.value] += 1
        return counts


def _atomic_write(path: Path, write_fn) -> None:
    """Escreve para um ficheiro temporário e só depois substitui o
    original — evita corromper o ficheiro de checkpoint se o processo
    for interrompido a meio da escrita."""
    fd, tmp_path = tempfile.mkstemp(dir=str(path.parent), suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="") as fh:
            write_fn(fh)
        os.replace(tmp_path, path)
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


class CsvDataSource(CheckpointDataSource):
    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        if not self.path.exists():
            raise FileNotFoundError(
                f"Ficheiro de dados não encontrado: {self.path}\n"
                "Veja data/sample_identifiers.csv como exemplo "
                f"(colunas: {COLUMN_IDENTIFICADOR}, {COLUMN_ESTADO}, {COLUMN_PROCESSADO_EM})."
            )

    def load_all(self) -> list[DataRecord]:
        with self.path.open("r", encoding="utf-8", newline="") as fh:
            reader = csv.DictReader(fh)
            if reader.fieldnames is None or COLUMN_IDENTIFICADOR not in reader.fieldnames:
                raise ValueError(
                    f"{self.path} precisa de uma coluna chamada '{COLUMN_IDENTIFICADOR}'."
                )
            records = []
            for row in reader:
                identificador = (row.get(COLUMN_IDENTIFICADOR) or "").strip()
                if not identificador:
                    continue
                records.append(
                    DataRecord(
                        identificador=identificador,
                        estado=RecordStatus.from_text(row.get(COLUMN_ESTADO)),
                        processado_em=(row.get(COLUMN_PROCESSADO_EM) or "").strip() or None,
                    )
                )
        return records

    def mark(self, identificador: str, estado: RecordStatus) -> None:
        records = self.load_all()
        timestamp = time.strftime("%Y-%m-%dT%H:%M:%S")
        found = False
        for record in records:
            if record.identificador == identificador and record.estado is RecordStatus.PENDENTE:
                record.estado = estado
                record.processado_em = timestamp
                found = True
                break
        if not found:
            # Se não havia nenhuma linha "pendente" com este identificador
            # (ex: já tinha sido marcada antes), atualiza a primeira que
            # corresponder de qualquer forma, para não perder o registo.
            for record in records:
                if record.identificador == identificador:
                    record.estado = estado
                    record.processado_em = timestamp
                    break

        def _write(fh):
            writer = csv.DictWriter(
                fh, fieldnames=[COLUMN_IDENTIFICADOR, COLUMN_ESTADO, COLUMN_PROCESSADO_EM]
            )
            writer.writeheader()
            for record in records:
                writer.writerow(
                    {
                        COLUMN_IDENTIFICADOR: record.identificador,
                        COLUMN_ESTADO: record.estado.value,
                        COLUMN_PROCESSADO_EM: record.processado_em or "",
                    }
                )

        _atomic_write(self.path, _write)


class ExcelDataSource(CheckpointDataSource):
    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        if not self.path.exists():
            raise FileNotFoundError(
                f"Ficheiro de dados não encontrado: {self.path}\n"
                f"Colunas esperadas: {COLUMN_IDENTIFICADOR}, {COLUMN_ESTADO}, {COLUMN_PROCESSADO_EM}."
            )

    def _load_workbook(self):
        import openpyxl

        return openpyxl.load_workbook(self.path)

    def _header_map(self, sheet) -> dict[str, int]:
        headers = {}
        for col_idx, cell in enumerate(sheet[1], start=1):
            if cell.value:
                headers[str(cell.value).strip().lower()] = col_idx
        if COLUMN_IDENTIFICADOR not in headers:
            raise ValueError(
                f"{self.path} precisa de uma coluna chamada '{COLUMN_IDENTIFICADOR}' na primeira linha."
            )
        return headers

    def load_all(self) -> list[DataRecord]:
        wb = self._load_workbook()
        sheet = wb.active
        headers = self._header_map(sheet)

        records = []
        for row in sheet.iter_rows(min_row=2):
            identificador_cell = row[headers[COLUMN_IDENTIFICADOR] - 1]
            identificador = str(identificador_cell.value or "").strip()
            if not identificador:
                continue
            estado_col = headers.get(COLUMN_ESTADO)
            processado_col = headers.get(COLUMN_PROCESSADO_EM)
            estado_val = row[estado_col - 1].value if estado_col else None
            processado_val = row[processado_col - 1].value if processado_col else None
            records.append(
                DataRecord(
                    identificador=identificador,
                    estado=RecordStatus.from_text(str(estado_val) if estado_val else None),
                    processado_em=str(processado_val) if processado_val else None,
                )
            )
        wb.close()
        return records

    def mark(self, identificador: str, estado: RecordStatus) -> None:
        wb = self._load_workbook()
        sheet = wb.active
        headers = self._header_map(sheet)

        # Garante que as colunas de checkpoint existem; se não existirem,
        # cria-as no fim da primeira linha.
        if COLUMN_ESTADO not in headers:
            new_col = sheet.max_column + 1
            sheet.cell(row=1, column=new_col, value=COLUMN_ESTADO)
            headers[COLUMN_ESTADO] = new_col
        if COLUMN_PROCESSADO_EM not in headers:
            new_col = sheet.max_column + 1
            sheet.cell(row=1, column=new_col, value=COLUMN_PROCESSADO_EM)
            headers[COLUMN_PROCESSADO_EM] = new_col

        timestamp = time.strftime("%Y-%m-%dT%H:%M:%S")
        id_col = headers[COLUMN_IDENTIFICADOR]
        estado_col = headers[COLUMN_ESTADO]
        processado_col = headers[COLUMN_PROCESSADO_EM]

        for row in sheet.iter_rows(min_row=2):
            cell_val = str(row[id_col - 1].value or "").strip()
            if cell_val == identificador:
                row[estado_col - 1].value = estado.value
                row[processado_col - 1].value = timestamp
                break

        wb.save(self.path)
        wb.close()


def open_data_source(path: str | Path) -> CheckpointDataSource:
    """Escolhe a implementação certa consoante a extensão do ficheiro."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return CsvDataSource(path)
    if suffix in (".xlsx", ".xlsm"):
        return ExcelDataSource(path)
    raise ValueError(
        f"Formato de dados não suportado: '{suffix}'. Use .csv ou .xlsx."
    )
